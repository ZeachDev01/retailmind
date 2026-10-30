from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

root = Path(r'C:/laragon/www/retailmind')
source = (root/'tmp/manual/build_manual.py').read_text(encoding='utf-8')
data = {}
exec(compile(source.split('\nstyles = {')[0], 'manual_content', 'exec'), data)
pages = data['pages']
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin, sec.bottom_margin = Pt(59), Pt(45)
sec.header_distance = Pt(25)

for name,size,line,before,after in [('Normal',11,14,0,5),('Title',26,33,0,18),('Subtitle',13,19,0,11),('Heading 1',17,21,0,13),('Heading 2',13,16,8,6),('Caption',9.5,12,5,10),('List Bullet',11,14,0,5)]:
    s=doc.styles[name]
    s.font.name='Times New Roman'
    s.font.size=Pt(size)
    s.font.color.rgb=RGBColor(0,0,0)
    fonts=s.element.get_or_add_rPr().find(qn('w:rFonts'))
    fonts.attrib.clear()
    for attr in ['ascii','hAnsi','eastAsia','cs']:
        fonts.set(qn('w:'+attr),'Times New Roman')
    for border in s.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    pf=s.paragraph_format
    pf.line_spacing=Pt(line)
    pf.space_before=Pt(before)
    pf.space_after=Pt(after)
    pf.widow_control=False
    if name in ['Heading 1','Heading 2']:
        pf.keep_with_next=True
    if name=='Title':
        s.font.bold=True
    if name=='Subtitle':
        s.font.italic=False
        s.font.bold=False
    if name=='Caption':
        s.font.italic=True
        s.font.bold=False

def field(p, instr, value=''):
    r=p.add_run()
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    command=OxmlElement('w:instrText'); command.set(qn('xml:space'),'preserve'); command.text=' '+instr+' '
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate')
    t=OxmlElement('w:t'); t.text=value
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    for x in [begin,command,sep,t,end]: r._r.append(x)

h=sec.header.paragraphs[0]
h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field(h,'PAGE','1')
def para(text='',style=None):
    return doc.add_paragraph(text,style)

p=para('RETAILMIND','Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(170)
p=para('USER MANUAL','Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('Shalom Store','Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10)
p=para('Inventory Sales and Demand Forecasting','Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('Version 1.0','Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(80)
p=para('31 October 2026','Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()
para('About this manual','Heading 1')
para('Use this manual to carry out daily Store tasks in RetailMind. It covers sign-in, platform governance, Store administration, stock management, checkout, GCash services, refunds, and recovery.')
para('Select the workspace assigned to your task. Cashiers handle checkout and shifts; Inventory Managers handle stock; Administrators oversee Store operations; the Super Administrator governs platform security and technical recovery.')
para('Screenshots show sample records. Follow the live page labels and validation messages when figures, available actions, or screen details differ.')
para('Contents','Heading 2')
for i,(title,_) in enumerate(pages):
    p=para(title+'\t')
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_after=Pt(2)
    field(p,f'PAGEREF topic_{i} \\h',str(i+3))
doc.add_page_break()

for i,(title,sections) in enumerate(pages):
    p=para(title,'Heading 1')
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(i)); start.set(qn('w:name'),f'topic_{i}')
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(i))
    p._p.insert(0,start); p._p.append(end)
    for s in sections:
        para(s['title'],'Heading 2')
        p=para(s['route']); p.paragraph_format.line_spacing=Pt(12); p.paragraph_format.space_after=Pt(6); p.paragraph_format.keep_with_next=True
        for r in p.runs: r.italic=True; r.font.size=Pt(10)
        for text in s['bullets']:
            p=para(text,'List Bullet')
            p.paragraph_format.left_indent=Pt(15)
            p.paragraph_format.first_line_indent=Pt(-10)
        if s['image']:
            path=root/s['image']
            with Image.open(path) as im: iw,ih=im.size
            width=min(468,205*iw/ih)
            p=para(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
            p.paragraph_format.keep_with_next=True
            pic=p.add_run().add_picture(str(path),width=Pt(width))
            pic._inline.docPr.set('descr',s['caption'])
            p=para(s['caption'],'Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if i<len(pages)-1: doc.add_page_break()

doc.core_properties.title='RetailMind User Manual'
doc.core_properties.author='RetailMind'
doc.core_properties.subject='User instructions for Shalom Store workspaces'
out=root/'output/doc/RetailMind_User_Manual.docx'
out.parent.mkdir(parents=True,exist_ok=True)
doc.save(out)
check=Document(out)
assert len(check.inline_shapes)==6
assert all(any(p.text==title for p in check.paragraphs) for title,_ in pages)
assert sum(len(s['bullets']) for _,ss in pages for s in ss)==sum(p.style.name=='List Bullet' for p in check.paragraphs)
print(out)
