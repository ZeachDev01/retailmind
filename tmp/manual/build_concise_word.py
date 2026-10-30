from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(r'C:/laragon/www/retailmind')
LIVE=ROOT/'tmp/manual/live'
captures=json.loads((LIVE/'operational-captures.json').read_text(encoding='utf-8'))
lookup={(x['role'],x['title']):x for x in captures}
screens=[]
def screen(role,title,*bullets):
    item=dict(lookup[(role,title)])
    item['bullets']=bullets
    screens.append(item)

role='Administrator'
screen(role,'Store Overview',
       'View Store performance, sales summaries, and operational alerts.',
       'Open dashboard links to review cash variances, pending Stock Issues, and Fiscal Periods.')
screen(role,'Store Staff',
       'View staff accounts, assigned roles, and account status.',
       'Add or update staff information and manage operational access.')
screen(role,'Registers',
       'View available Registers and their current status.',
       'Manage Register names, receipt paper widths, and availability.')
screen(role,'Store Settings',
       'View operational thresholds and review limits for the Store.',
       'Update the required values and select Save settings.')
screen(role,'Store Calendar',
       'View the Store calendar and daily operating statuses.',
       'Record open days, closures, holidays, and exceptional closures.')
screen(role,'Fiscal Periods',
       'View and create Fiscal Periods for Store records.',
       'Review outstanding responsibilities before closing or locking a period.')
screen(role,'Promotions',
       'View and manage Store discounts and promotions.',
       'Set the promotion scope, discount, minimum quantity, and active dates.')
screen(role,'GCash Overview',
       'View recorded service cash, wallet balances, and GCash activity.',
       'Review transactions, closing checks, and unresolved exchanges.')
screen(role,'Cashier Shifts',
       'View Cashier Shift records, sales, drawer counts, and cash differences.',
       'Filter by date, Cashier, or Register to review the relevant shift.')
screen(role,'Database Backup',
       'Create and download an encrypted copy of Store records.',
       'Use a passphrase of at least 16 characters and keep it separately; lost passphrases cannot be recovered.')
screen(role,'Generate Report',
       'View sales and other available Store reports.',
       'Choose the report and filters, then run, export, or print the results.')
screen(role,'Operational Audit',
       'View recorded Store actions, users, modules, and timestamps.',
       'Search the history and open details to trace an operational event.')
screen(role,'Receipt History',
       'View completed Store receipts and related refund information.',
       'Filter the list and open Full details to review the selected receipt.')
screen(role,'Refund Correction Reviews',
       'Review corrections needed after a recorded refund.',
       'Approve the exact cash or stock changes and assign the responsible staff.')

role='Inventory Manager'
screen(role,'Stock Overview',
       'View stock totals, low-stock alerts, and expiry information.',
       'Use the overview to identify products that need attention.')
screen(role,'Products',
       'View product names, prices, stock levels, and status.',
       'Add or update products and manage their barcode and packaging details.')
screen(role,'Replenishment Recommendations',
       'View suggested buying quantities based on stock rules and available forecasts.',
       'Use recommendations as buying guidance; an unavailable forecast does not mean zero demand.')
screen(role,'Stock Receiving',
       'Record bought goods using the product catalog or barcode search.',
       'Review quantities, costs, and batch details before saving accepted stock.')
screen(role,'Receiving History',
       'View the stock receipts available to your account.',
       'Search recorded entries and open their details to review received quantities.')
screen(role,'Inventory Counts',
       'Record physical product or batch counts and the reason for a discrepancy.',
       'Review pending counts before approval; stock changes only after approval.')
screen(role,'Stock Issue Reviews',
       'View Stock Issues submitted for Inventory Manager review.',
       'Check the reported goods and record the appropriate decision and notes.')
screen(role,'Supplier Pullouts',
       'Record goods removed from a selected product batch.',
       'Enter the actual quantity and reason, then review the pullout history.')
screen(role,'Demand Forecast',
       'View available product demand estimates and forecast status.',
       'Check the forecast period and readiness before using estimates for planning.')
screen(role,'Forecast Exceptions',
       'View products whose forecast results need further review.',
       'Use the exception filters to identify the affected products and reasons.')
screen(role,'Data Readiness',
       'View whether product history is ready for forecasting.',
       'Identify products needing more history or further data review.')
screen(role,'Insights & Risk',
       'View inventory risks, including excess stock, expiry exposure, and slow-moving goods.',
       'Use the analysis and cycle-count information to prioritize stock checks.')
screen(role,'Generate Report',
       'View the reports available in the Inventory workspace.',
       'Select the report and filters, then review or export its results.')
screen(role,'Refund Correction Reviews',
       'View stock correction tasks assigned after an Administrator review.',
       'Record the approved stock action and notes while preserving the original refund.')

role='Cashier'
screen(role,'Cashier Shift',
       'Open a shift using an available Register and the actual starting cash.',
       'Review drawer movements and count cash before closing the shift.')
screen(role,'Point of Sale',
       'Scan or search products, review quantities, and complete customer sales.',
       'Open a Cashier Shift before checkout and print the receipt after a successful sale.')
screen(role,'GCash',
       'Record Cash-In and Cash-Out exchanges, fees, and reference numbers.',
       'Verify actual wallet and cash movements, then reconcile or hand over service funds.')
screen(role,'Dashboard Overview',
       'View daily sales, transaction totals, stock alerts, and recent activity.',
       'Use the Cashier and GCash tabs to review the corresponding summaries.')
screen(role,'Cash Refunds',
       'Find the original sale and record returned quantities, item condition, and the refund reason.',
       'Only Restockable items return to available stock; complete the required payment verification and approvals.')
screen(role,'Transaction History',
       'View your sales, Cash Refunds, drawer movements, and shifts.',
       'Choose a history tab and date range to find the relevant record.')
screen(role,'Receipt History',
       'View the receipts available to your Cashier account.',
       'Open a saved receipt to review or print it without repeating the sale.')
screen(role,'Report Stock Issue',
       'Report damaged, missing, expired, or other affected stock during an open shift.',
       'Identify the product, stock form, quantity, and explanation for Inventory Manager review.')
screen(role,'Refund Correction Reviews',
       'View cash correction tasks assigned after an Administrator review.',
       'Record the approved drawer action and required notes without rewriting the original refund.')
screen(role,'Notifications',
       'View notifications and unread items for your account.',
       'Open the related record to review any action that needs attention.')

assert len(screens)==38
assert all(len(x['bullets'])==2 for x in screens)
assert all((LIVE/x['file']).is_file() for x in screens)

doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.left_margin=section.right_margin=Inches(1)
section.top_margin=Pt(58); section.bottom_margin=Pt(45)
section.header_distance=Pt(25)
for name,size,leading in [('Normal',12,16),('Title',24,30),('Subtitle',13,19),('Heading 1',12,16),('List Bullet',12,16)]:
    style=doc.styles[name]
    style.font.name='Times New Roman'
    style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.italic=False
    fonts=style.element.get_or_add_rPr().find(qn('w:rFonts'))
    fonts.attrib.clear()
    for attr in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+attr),'Times New Roman')
    for border in style.element.xpath('.//w:pBdr'): border.getparent().remove(border)
    pf=style.paragraph_format
    pf.line_spacing=Pt(leading); pf.space_before=Pt(0); pf.space_after=Pt(5)
    pf.widow_control=False
doc.styles['Heading 1'].font.bold=True
doc.styles['Title'].font.bold=True
doc.styles['Subtitle'].font.bold=False

header=section.header.paragraphs[0]
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
header._p.append(field)

p=doc.add_paragraph('USER MANUAL','Title')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before=Pt(190)
p.paragraph_format.space_after=Pt(24)
p=doc.add_paragraph('RetailMind','Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Administrator  Inventory Manager  Cashier','Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

for index,item in enumerate(screens):
    heading=item['title'] if item['title'].startswith(item['role']+' ') else item['role']+' '+item['title']
    p=doc.add_paragraph(heading,'Heading 1')
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_after=Pt(9)
    for text in item['bullets']:
        p=doc.add_paragraph(text,'List Bullet')
        p.paragraph_format.left_indent=Pt(20)
        p.paragraph_format.first_line_indent=Pt(-12)
        p.paragraph_format.keep_with_next=True
    file=LIVE/item['file']
    with Image.open(file) as im: iw,ih=im.size
    width=min(468,230*iw/ih)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1
    image=p.add_run().add_picture(str(file),width=Pt(width))
    image._inline.docPr.set('descr',f"Live RetailMind {item['role']} {item['title']} screen")
    if index%2==0:
        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.line_spacing=Pt(18)
        p.add_run(' ').font.size=Pt(1)
    elif index<len(screens)-1:
        doc.add_page_break()

doc.core_properties.title='RetailMind User Manual'
doc.core_properties.author='RetailMind'
doc.core_properties.subject='Administrator Inventory Manager and Cashier overview'
out=ROOT/'output/doc/RetailMind_User_Manual_Admin_Inventory_Cashier.docx'
doc.save(out)
check=Document(out)
assert len(check.inline_shapes)==38
assert sum(p.style.name=='List Bullet' for p in check.paragraphs)==76
assert not check.styles['Title'].element.xpath('.//w:pBdr')
(LIVE/'manual-sections.json').write_text(json.dumps([{k:x[k] for k in ['role','title','file','bullets']} for x in screens],indent=2),encoding='utf-8')
print(out)
